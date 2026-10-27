"""
Exportador Excel — 4 tabs estandar:
  1. Catalogo        — todos los productos, todos los campos
  2. Por Categoria   — agrupado por categoria + linea
  3. Por Marca       — agrupado por marca/fabricante
  4. Resumen         — estadisticas generales
"""

import logging
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("guadalajara")

# Colores brand Farmacias Guadalajara (verde oscuro + blanco)
COLOR_HEADER  = "1B6B2F"   # verde oscuro FGL
COLOR_ACCENT  = "E8F5E9"   # verde muy claro (filas alternas)
COLOR_WHITE   = "FFFFFF"
COLOR_WARN    = "FFF9C4"   # amarillo claro (sin EAN)


# ──────────────────────────────────────────────────────────────
# COLUMNAS DEL CATALOGO
# ──────────────────────────────────────────────────────────────
CATALOG_COLUMNS = [
    # (key en el dict, header visible, ancho)
    ("pid",                "PID",             10),
    ("ean",                "EAN",             16),
    ("nombre",             "Nombre",          50),
    ("marca",              "Marca",           20),
    ("fabricante",         "Fabricante",      35),
    ("categoria",          "Categoría",       22),
    ("linea",              "Línea",           28),
    ("precio_venta",       "Precio Venta",    14),
    ("precio_lista",       "Precio Lista",    14),
    ("descuento_pct",      "Descuento %",     13),
    ("disponible",         "Disponible",      12),
    ("disponible_tienda",  "Disp. Tienda",    13),
    ("stock_online",       "Stock Online",    13),
    ("es_antibiotico",     "Antibiótico",     13),
    ("es_controlado",      "Controlado",      12),
    ("refrigerar",         "Refrigerar",      12),
    ("marca_propia",       "Marca Propia",    13),
    ("es_cbd",             "CBD",             10),
    ("iva_pct",            "IVA %",           10),
    ("ieps_pct",           "IEPS %",          10),
    ("es_a_granel",        "A Granel",        12),
    ("regional",           "Regional",        10),
    ("promociones",        "Promociones",     25),
    ("plm_id",             "PLM ID",          10),
    ("descripcion_interna","Desc. Interna",   30),
    ("descripcion_larga",  "Descripción",     50),
    ("principio_activo",   "Principio Activo",50),
    ("url_imagen",         "URL Imagen",      60),
    ("url",                "URL Producto",    70),
]


def export_to_excel(products, output_path):
    """Crea el archivo .xlsx con las 4 tabs."""
    logger.info(f"Exportando {len(products)} productos a Excel...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # eliminar hoja default

    _write_catalogo(wb, products)
    _write_por_categoria(wb, products)
    _write_por_marca(wb, products)
    _write_resumen(wb, products)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Excel guardado: {output_path}")


# ──────────────────────────────────────────────────────────────
# TAB 1: CATALOGO
# ──────────────────────────────────────────────────────────────
def _write_catalogo(wb, products):
    ws = wb.create_sheet("Catálogo")

    header_fill  = PatternFill("solid", fgColor=COLOR_HEADER)
    header_font  = Font(bold=True, color=COLOR_WHITE, size=10)
    accent_fill  = PatternFill("solid", fgColor=COLOR_ACCENT)
    warn_fill    = PatternFill("solid", fgColor=COLOR_WARN)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Cabecera
    headers = [col[1] for col in CATALOG_COLUMNS]
    widths  = [col[2] for col in CATALOG_COLUMNS]

    for col_idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font   = header_font
        cell.fill   = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # Filas de datos
    for row_idx, prod in enumerate(products, start=2):
        is_alt = (row_idx % 2 == 0)
        has_ean = bool(prod.get("ean"))
        row_fill = warn_fill if not has_ean else (accent_fill if is_alt else None)

        for col_idx, (key, header, width) in enumerate(CATALOG_COLUMNS, start=1):
            val = prod.get(key, "")
            # Formatear booleanos
            if isinstance(val, bool):
                val = "Sí" if val else "No"
            # Formatear None
            if val is None:
                val = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            if row_fill:
                cell.fill = row_fill

            # Alineacion por tipo
            if key in ("precio_venta", "precio_lista", "descuento_pct", "iva_pct", "ieps_pct"):
                cell.alignment = Alignment(horizontal="right")
                if isinstance(val, float):
                    cell.number_format = "#,##0.00"
            elif key in ("ean", "pid", "plm_id"):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(wrap_text=False)

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions


# ──────────────────────────────────────────────────────────────
# TAB 2: POR CATEGORIA
# ──────────────────────────────────────────────────────────────
def _write_por_categoria(wb, products):
    ws = wb.create_sheet("Por Categoría")

    headers = [
        "Categoría", "Línea", "# Productos", "Con EAN",
        "% EAN", "Precio Prom.", "Precio Min.", "Precio Max.",
        "Antibióticos", "Controlados", "Refrigerados"
    ]
    widths  = [22, 28, 13, 10, 10, 14, 12, 12, 13, 12, 13]

    _write_header_row(ws, headers, widths)
    ws.freeze_panes = "A2"

    # Agrupar
    groups = defaultdict(list)
    for p in products:
        key = (p.get("categoria", "") or "SIN CATEGORIA",
               p.get("linea", "") or "")
        groups[key].append(p)

    header_fill = PatternFill("solid", fgColor=COLOR_HEADER)
    accent_fill = PatternFill("solid", fgColor=COLOR_ACCENT)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_font = Font(bold=True, color=COLOR_WHITE, size=10)

    row = 2
    for (cat, linea), prods in sorted(groups.items()):
        prices    = [p.get("precio_venta") for p in prods if p.get("precio_venta") is not None]
        with_ean  = sum(1 for p in prods if p.get("ean"))
        antib     = sum(1 for p in prods if p.get("es_antibiotico"))
        control   = sum(1 for p in prods if p.get("es_controlado"))
        refrig    = sum(1 for p in prods if p.get("refrigerar"))

        vals = [
            cat, linea, len(prods), with_ean,
            round(with_ean / len(prods) * 100, 1) if prods else 0,
            round(sum(prices) / len(prices), 2) if prices else None,
            min(prices) if prices else None,
            max(prices) if prices else None,
            antib, control, refrig,
        ]

        fill = accent_fill if row % 2 == 0 else None
        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = border
            if fill:
                cell.fill = fill
        row += 1

    ws.auto_filter.ref = ws.dimensions


# ──────────────────────────────────────────────────────────────
# TAB 3: POR MARCA
# ──────────────────────────────────────────────────────────────
def _write_por_marca(wb, products):
    ws = wb.create_sheet("Por Marca")

    headers = [
        "Marca", "Fabricante", "# Productos", "Con EAN",
        "% EAN", "Precio Prom.", "Precio Min.", "Precio Max.", "Marca Propia"
    ]
    widths  = [22, 35, 13, 10, 10, 14, 12, 12, 13]

    _write_header_row(ws, headers, widths)
    ws.freeze_panes = "A2"

    groups = defaultdict(list)
    for p in products:
        key = (p.get("marca", "") or "SIN MARCA",
               p.get("fabricante", "") or "")
        groups[key].append(p)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    accent_fill = PatternFill("solid", fgColor=COLOR_ACCENT)

    row = 2
    for (marca, fab), prods in sorted(groups.items()):
        prices   = [p.get("precio_venta") for p in prods if p.get("precio_venta") is not None]
        with_ean = sum(1 for p in prods if p.get("ean"))
        mp       = sum(1 for p in prods if p.get("marca_propia"))

        vals = [
            marca, fab, len(prods), with_ean,
            round(with_ean / len(prods) * 100, 1) if prods else 0,
            round(sum(prices) / len(prices), 2) if prices else None,
            min(prices) if prices else None,
            max(prices) if prices else None,
            mp,
        ]

        fill = accent_fill if row % 2 == 0 else None
        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = border
            if fill:
                cell.fill = fill
        row += 1

    ws.auto_filter.ref = ws.dimensions


# ──────────────────────────────────────────────────────────────
# TAB 4: RESUMEN
# ──────────────────────────────────────────────────────────────
def _write_resumen(wb, products):
    ws = wb.create_sheet("Resumen")

    header_fill = PatternFill("solid", fgColor=COLOR_HEADER)
    header_font = Font(bold=True, color=COLOR_WHITE, size=11)
    label_font  = Font(bold=True, size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 22

    total     = len(products)
    with_ean  = sum(1 for p in products if p.get("ean"))
    prices    = [p.get("precio_venta") for p in products if p.get("precio_venta") is not None]
    cats      = len(set(p.get("categoria","") for p in products))
    marcas    = len(set(p.get("marca","") for p in products))
    with_disc = sum(1 for p in products if (p.get("descuento_pct") or 0) > 0)
    antib     = sum(1 for p in products if p.get("es_antibiotico"))
    control   = sum(1 for p in products if p.get("es_controlado"))
    refrig    = sum(1 for p in products if p.get("refrigerar"))
    disp      = sum(1 for p in products if p.get("disponible"))
    mp        = sum(1 for p in products if p.get("marca_propia"))

    rows_data = [
        ("FARMACIAS GUADALAJARA — RESUMEN", None),
        (None, None),
        ("CATALOGO", None),
        ("Total productos",     total),
        ("Categorías únicas",   cats),
        ("Marcas únicas",       marcas),
        (None, None),
        ("EAN / CODIGO DE BARRAS", None),
        ("Productos con EAN",   with_ean),
        ("Sin EAN",             total - with_ean),
        ("Cobertura EAN",       f"{with_ean/total*100:.1f}%" if total else "0%"),
        (None, None),
        ("PRECIOS", None),
        ("Precio promedio",     f"${sum(prices)/len(prices):.2f}" if prices else "N/A"),
        ("Precio mínimo",       f"${min(prices):.2f}" if prices else "N/A"),
        ("Precio máximo",       f"${max(prices):.2f}" if prices else "N/A"),
        ("Con descuento",       with_disc),
        (None, None),
        ("TIPO DE PRODUCTO", None),
        ("Antibióticos",        antib),
        ("Controlados",         control),
        ("Requieren refrigeración", refrig),
        ("Marca propia",        mp),
        ("Disponibles",         disp),
    ]

    for row_idx, (label, value) in enumerate(rows_data, start=1):
        if label is None:
            continue

        cell_a = ws.cell(row=row_idx, column=1, value=label)
        cell_b = ws.cell(row=row_idx, column=2, value=value)

        # Secciones de cabecera
        if value is None and label:
            cell_a.font = header_font
            cell_a.fill = header_fill
            cell_b.fill = header_fill
        else:
            cell_a.font = label_font

        cell_a.border = border
        cell_b.border = border
        cell_b.alignment = Alignment(horizontal="center")


# ──────────────────────────────────────────────────────────────
# HELPER
# ──────────────────────────────────────────────────────────────
def _write_header_row(ws, headers, widths):
    header_fill = PatternFill("solid", fgColor=COLOR_HEADER)
    header_font = Font(bold=True, color=COLOR_WHITE, size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.border    = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
